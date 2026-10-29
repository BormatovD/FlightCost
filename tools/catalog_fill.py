"""Дозаполнить каталог типов ВС из таблицы Excel.

Зачем. Тип попадает в расчёт маршрута, только если у него есть паспорт
(`REQUIRED` в parsers/aircraft_user.py) и расход: вычислитель openap или
аналог с якорем — опубликованный рейс «дистанция · топливо · пассажиры»
(решения 105, 224). У 24 типов каталога чего-то из этого нет, и они только
занимают место в списке. Таблица — способ собрать недостающее руками и
одним шагом переложить в `data/aircraft/catalog/*.json`.

    python tools/catalog_fill.py template типы-дозаполнить.xlsx
        таблица с недостающим: лист «Паспорт» — по строке на пустое поле,
        лист «Расход» — по строке на тип без якоря (у турбовинтовых — три
        дистанции), лист «Как заполнять».

    python tools/catalog_fill.py apply типы-дозаполнить.xlsx [--dry-run]
        заполненные строки → JSON каталога. Каждый изменённый файл
        проверяется тем же парсером, что `fca refresh`; файл, который
        парсер не принял, не пишется. Дальше как обычно:
        fca refresh --only aircraft_catalog → fca review → принять.

Вид данных в каждой строке: «документ» (опубликовано издателем — адрес
обязателен) или «оценка разработчика» (собрано нами из нескольких
источников, пересчитано, оценено). Оценка помечается в каталоге меткой
`seo.DEV_PREFIX` — в `meta.assumed` для поля паспорта и в начале
`anchor_source` для расхода, — и страница типа на сайте пишет это словами.

Турбовинтовых аналогов в openap нет: якорь у ATR, Dash 8, L-410 и других
сейчас можно привязать только к реактивному вычислителю, и время рейса и
расход на других дистанциях будут неверны. Поэтому для них таблица просит
три дистанции: строки сохраняются в `burn_table` и станут основой
табличного вычислителя; аналог для них лучше не называть.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from flightcostapp.parsers.aircraft_user import FIELDS, REQUIRED, _one  # noqa: E402
from flightcostapp.seo import DEV_PREFIX                                  # noqa: E402

CATALOG = ROOT / "data" / "aircraft" / "catalog"
KINDS = ("документ", "оценка разработчика")
LABEL = {"mtow_t": "MTOW, т", "oew_t": "Масса пустого снаряжённого (OEW), т",
         "fuel_capacity_t": "Запас топлива, т", "seats_max": "Кресел, максимум по сертификату",
         "span_m": "Размах крыла, м", "cruise_mach": "Крейсерское число M",
         "cruise_alt_m": "Крейсерская высота, м", "mlw_t": "Максимальная посадочная масса, т",
         "seats_typical": "Кресел, типовая", "length_m": "Длина, м", "engine_count": "Число двигателей"}
P_HEAD = ["ИКАО", "Тип", "Поле", "Что это", "Сейчас", "Значение", "Источник (издатель, документ, год)",
          "Адрес", "Вид данных", "Примечание"]
B_HEAD = ["ИКАО", "Тип", "Аналог (openap)", "Дистанция, nm", "Блок-время, мин", "Топливо на рейс, кг",
          "Пассажиров", "Источник (издатель, документ, год)", "Адрес", "Вид данных", "Примечание"]


def _load() -> dict[str, tuple[Path, dict]]:
    out = {}
    for p in sorted(CATALOG.glob("*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        icao = str(d.get("icao") or "").upper()
        if icao and not icao.startswith("USER:"):
            out[icao] = (p, d)
    return out


def _needs(d: dict) -> tuple[list[str], str]:
    """(пустые обязательные поля, что с расходом: ok | anchor | table)."""
    f = d.get("fields") or {}
    missing = [k for k in REQUIRED if f.get(k) in (None, "")]
    b = d.get("burn") or {}
    if b.get("same_engine") or b.get("anchor_kg_per_h") is not None or b.get("anchor_trip_kg") is not None:
        burn = "ok"
    else:
        mach = f.get("cruise_mach")
        burn = "table" if isinstance(mach, (int, float)) and mach < 0.65 else "anchor"
    return missing, burn


def template(path: Path) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="262626")
    known = PatternFill("solid", fgColor="EDEDED")
    todo = PatternFill("solid", fgColor="FFF4D6")

    def sheet(ws, cols, widths):
        ws.append(cols)
        for c, w in zip(ws[1], widths):
            c.font, c.fill = head, hfill
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[c.column_letter].width = w
        ws.freeze_panes = "A2"
        dv = DataValidation(type="list", formula1='"' + ",".join(KINDS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        return dv

    ws = wb.active
    ws.title = "Паспорт"
    dv_p = sheet(ws, P_HEAD, [8, 30, 16, 34, 10, 12, 40, 40, 20, 30])
    ws2 = wb.create_sheet("Расход")
    dv_b = sheet(ws2, B_HEAD, [8, 30, 14, 13, 14, 16, 12, 40, 40, 20, 30])
    n = 0
    for icao, (_p, d) in _load().items():
        missing, burn = _needs(d)
        name = d.get("name") or icao
        cur = d.get("fields") or {}
        for k in missing:
            ws.append([icao, name, k, LABEL.get(k, k), cur.get(k), None, None, None, None, None])
            n += 1
            r = ws.max_row
            for c in ws[r][:5]:
                c.fill = known
            for c in ws[r][5:9]:
                c.fill = todo
            dv_p.add(f"I{r}")
        if burn == "ok":
            continue
        analog = (d.get("burn") or {}).get("analog") or ""
        dists = (200, 400, 600) if burn == "table" else (None,)
        note = ("турбовинтовой: три дистанции, аналог не называть" if burn == "table"
                else "один опубликованный рейс; дистанция — типичная для типа")
        for nm in dists:
            ws2.append([icao, name, analog, nm, None, None, None, None, None, None, note])
            n += 1
            r = ws2.max_row
            for c in ws2[r][:3]:
                c.fill = known
            for c in ws2[r][3:10]:
                c.fill = todo
            dv_b.add(f"J{r}")
    ws3 = wb.create_sheet("Как заполнять")
    for line in (
        "Серые ячейки — что уже есть в каталоге; жёлтые — заполнить. Пустую строку можно оставить: она пропускается.",
        "Вид данных: «документ» — число опубликовано издателем (тогда «Адрес» обязателен); «оценка разработчика» — "
        "собрано, пересчитано или оценено нами. Оценка будет так и подписана на сайте.",
        "Паспорт: одно значение на строку, в единицах из столбца «Что это». Можно добавить строку для любого поля "
        f"из списка: {', '.join(FIELDS)}.",
        "Расход: рейс целиком — дистанция по маршруту в морских милях, топливо на рейс (block fuel) в кг, число "
        "пассажиров на этом рейсе. Галлоны × 3,039 = кг; км ÷ 1,852 = nm.",
        "Реактивные типы: хватит одной строки — расчёт масштабирует вычислитель аналога под этот рейс. Аналог можно "
        "поменять на любой тип openap с моделью расхода (A20N, A320, A321, B38M, B738, E75L, E190, E195, B752, B788 …).",
        "Турбовинтовые: три дистанции (200 / 400 / 600 nm или ближайшие опубликованные). Аналога в openap нет; "
        "строки сохраняются для табличного вычислителя, в расчёт тип войдёт, когда он будет готов.",
        "Применить: python tools/catalog_fill.py apply <файл>.xlsx, затем fca refresh --only aircraft_catalog, "
        "fca review, принять. JSON каталога закоммитить.",
    ):
        ws3.append([line])
    ws3.column_dimensions["A"].width = 140
    for row in ws3.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    print(f"{path}: строк к заполнению {n}")
    return n


def _rows(ws, head):
    it = ws.iter_rows(values_only=True)
    cols = next(it)
    if list(cols[:len(head)]) != head:
        raise ValueError(f"лист «{ws.title}»: заголовок не тот, что у шаблона — сначала "
                         f"python tools/catalog_fill.py template")
    for r in it:
        yield dict(zip(head, r))


def _kind(v) -> str:
    k = str(v or "").strip().lower()
    if k not in KINDS:
        raise ValueError(f"вид данных {v!r}: «документ» или «оценка разработчика»")
    return k


def apply(path: Path, dry: bool = False) -> dict[str, list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    cat = _load()
    changed: dict[str, list[str]] = {}
    errors: list[str] = []
    today = date.today().isoformat()

    for i, r in enumerate(_rows(wb["Паспорт"], P_HEAD), start=2):
        if r["Значение"] in (None, ""):
            continue
        icao, k = str(r["ИКАО"] or "").upper(), str(r["Поле"] or "").strip()
        try:
            if icao not in cat:
                raise ValueError(f"типа {icao} нет в каталоге")
            if k not in FIELDS:
                raise ValueError(f"поле {k!r} не из списка паспорта")
            kind, src, url = _kind(r["Вид данных"]), str(r["Источник (издатель, документ, год)"] or "").strip(), \
                str(r["Адрес"] or "").strip()
            if not src:
                raise ValueError("нет источника")
            if kind == "документ" and not url.startswith("http"):
                raise ValueError("для документа нужен адрес")
            d = cat[icao][1]
            d.setdefault("fields", {})[k] = float(r["Значение"])
            meta = d.setdefault("meta", {})
            meta.setdefault("field_sources", {})[k] = f"{src} {url}".strip()
            if kind == "оценка разработчика":
                meta.setdefault("assumed", {})[k] = f"{DEV_PREFIX}: {src}" + (f" {url}" if url else "")
            else:
                (meta.get("assumed") or {}).pop(k, None)
            changed.setdefault(icao, []).append(f"{k} = {r['Значение']} ({kind})")
        except (ValueError, TypeError) as exc:
            errors.append(f"Паспорт, строка {i}: {exc}")

    by_type: dict[str, list[dict]] = {}
    for i, r in enumerate(_rows(wb["Расход"], B_HEAD), start=2):
        if r["Топливо на рейс, кг"] in (None, ""):
            continue
        r["_row"] = i
        by_type.setdefault(str(r["ИКАО"] or "").upper(), []).append(r)
    for icao, rows in by_type.items():
        try:
            if icao not in cat:
                raise ValueError(f"типа {icao} нет в каталоге")
            d = cat[icao][1]
            table = []
            for r in rows:
                kind = _kind(r["Вид данных"])
                src, url = str(r["Источник (издатель, документ, год)"] or "").strip(), str(r["Адрес"] or "").strip()
                if not src:
                    raise ValueError(f"строка {r['_row']}: нет источника")
                if kind == "документ" and not url.startswith("http"):
                    raise ValueError(f"строка {r['_row']}: для документа нужен адрес")
                if r["Дистанция, nm"] in (None, "") or r["Пассажиров"] in (None, ""):
                    raise ValueError(f"строка {r['_row']}: нужны дистанция и пассажиры вместе с топливом")
                table.append({"nm": float(r["Дистанция, nm"]), "kg": float(r["Топливо на рейс, кг"]),
                              "pax": float(r["Пассажиров"]),
                              "block_min": float(r["Блок-время, мин"]) if r["Блок-время, мин"] else None,
                              "kind": kind, "source": f"{src} {url}".strip(),
                              "note": str(r["Примечание"] or "").strip()})
            # Все рейсы — в `burn_table` рядом с паспортом (парсер её не читает):
            # основа будущего табличного вычислителя и проверка якоря.
            d["burn_table"] = sorted(table, key=lambda x: x["nm"])
            analog = str(rows[0]["Аналог (openap)"] or (d.get("burn") or {}).get("analog") or "").strip().upper()
            if analog:
                # Якорь — рейс, ближайший к середине таблицы: одна точка, по ней
                # масштабируется вычислитель аналога (решение 224).
                b = d.setdefault("burn", {})
                mid = d["burn_table"][len(d["burn_table"]) // 2]
                b["analog"] = analog
                b["anchor_trip_kg"], b["anchor_trip_nm"], b["anchor_trip_pax"] = mid["kg"], mid["nm"], mid["pax"]
                b.pop("anchor_kg_per_h", None)
                b.pop("anchor_mass_t", None)
                pre = f"{DEV_PREFIX}: " if mid["kind"] == "оценка разработчика" else ""
                b["anchor_source"] = pre + mid["source"]
                changed.setdefault(icao, []).append(f"якорь {mid['nm']:.0f} nm · {mid['kg']:.0f} кг · "
                                                    f"{mid['pax']:.0f} пасс. ({mid['kind']}), аналог {analog}")
            else:
                changed.setdefault(icao, []).append(f"таблица расхода, {len(table)} строк (без аналога — "
                                                    f"ждёт табличного вычислителя)")
        except (ValueError, TypeError) as exc:
            errors.append(f"Расход, {icao}: {exc}")

    for icao in list(changed):
        p, d = cat[icao]
        d["checked_at"] = today
        try:
            _one(json.loads(json.dumps(d)), {"source_id": "aircraft_catalog"})   # разбор как у fca refresh
        except Exception as exc:                              # noqa: BLE001
            errors.append(f"{icao}: парсер каталога не принял — {exc}")
            changed.pop(icao)
            continue
        if not dry:
            p.write_text(json.dumps(d, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    for icao, what in changed.items():
        print(f"{icao}: " + "; ".join(what))
    for e in errors:
        print("ОШИБКА " + e, file=sys.stderr)
    if not changed and not errors:
        print("заполненных строк нет")
    if errors:
        raise SystemExit(1)
    return changed


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("template")
    t.add_argument("xlsx", type=Path)
    a = sub.add_parser("apply")
    a.add_argument("xlsx", type=Path)
    a.add_argument("--dry-run", action="store_true")
    ns = ap.parse_args(argv)
    if ns.cmd == "template":
        template(ns.xlsx)
    else:
        apply(ns.xlsx, ns.dry_run)


if __name__ == "__main__":
    main()
