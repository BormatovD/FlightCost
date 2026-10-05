"""Разбор базы шума EASA — на выдержке из НАСТОЯЩЕГО файла.

Первая версия парсера проверялась на книге с выдуманными заголовками и
была зелёной. На настоящем файле шапка оказалась в три строки с
объединёнными ячейками, и парсер не нашёл ни числа двигателей, ни
двигателя, ни пределов: четырёхдвигательные типы получили предел пролёта
для двух двигателей. Поэтому фикстура — настоящие строки с настоящей
шапкой (выпуск 53 от 26.06.2026, источник указан на листе SOURCE), а
проверяются свойства, которые обязаны держаться в любой редакции.

Сеть не нужна: файл в `tests/fixtures/`. Нужен `openap` — он даёт MTOW и
двигатель типа для выбора строки, как и в продукте.
"""

from __future__ import annotations

import pathlib
import unittest

from flightcostapp.categories import chapter3_limits
from flightcostapp.parsers import easa_noise as E

FIX = pathlib.Path(__file__).parent / "fixtures" / "easa_noise_excerpt.xlsx"


class EasaNoise(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.blob = FIX.read_bytes()
        cls.ctx = {"source_id": "easa_noise", "today": "2026-10-01", "min_types": 5}
        cls.facts = E.parse(cls.blob, cls.ctx)
        cls.by = {f.key: f for f in cls.facts}

    def v(self, key):
        self.assertIn(key, self.by, f"нет факта {key}")
        return self.by[key].value

    def test_multirow_header_is_read_whole(self):
        """Все столбцы найдены, включая разбитые на несколько строк шапки."""
        import io
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(self.blob), read_only=True,
                                    data_only=True)["JETS"]
        _, col, _ = E._find_header(list(ws.iter_rows(values_only=True)))
        for need in ("engines", "engine", "lim_lat", "lim_fly", "lim_app",
                     "mg_cum", "chapter"):
            self.assertIn(need, col, f"столбец {need} не найден в шапке EASA")

    def test_four_engines_counted(self):
        """Четырёхдвигательный тип получает предел пролёта для четырёх."""
        self.assertEqual(self.v("B744/engine_count_cert"), 4)
        self.assertEqual(self.v("A388/engine_count_cert"), 4)
        self.assertEqual(self.v("A320/engine_count_cert"), 2)

    def test_our_chapter3_formula_matches_publisher(self):
        """Вторая реализация той же величины — только с перекрёстной
        проверкой (решение 102): ни одного расхождения в отчёте."""
        bad = [s for s in self.ctx.get("summary", []) if "расходится" in s
               or "по нашей формуле" in s]
        self.assertEqual(bad, [], "наша формула главы 3 разошлась с файлом")

    def test_neighbours_not_confused(self):
        """737-8 — не 737-800, 747-400F — не 747-400, 777-300ER — не 777-300."""
        self.assertIn("737-8", self.by["B38M/noise_mtom_t"].note)
        self.assertNotIn("737-800", self.by["B38M/noise_mtom_t"].note)
        self.assertIn("737-800", self.by["B738/noise_mtom_t"].note)
        self.assertNotIn("747-400F", self.by["B744/noise_mtom_t"].note)
        self.assertIn("777-300ER", self.by["B77W/noise_mtom_t"].note)
        self.assertNotIn("777-300ER", self.by["B773/noise_mtom_t"].note)

    def test_neo_quieter_than_ceo(self):
        """Если neo не тише ceo — выбрана не та строка."""
        self.assertGreater(self.v("A20N/noise_margin_cum_epndb"),
                           self.v("A320/noise_margin_cum_epndb"))

    def test_engine_of_the_type_is_chosen(self):
        """Строка выбирается по двигателю типа, а не только по массе."""
        self.assertNotIn("двигатель типа не найден", self.by["A320/noise_mtom_t"].note)

    def test_chapter3_limits_join_at_band_edges(self):
        for m, e in ((35.0, 2), (48.1, 2), (280, 2), (385, 2), (400, 2),
                     (28.6, 3), (20.2, 4)):
            lo, hi = chapter3_limits(m - 1e-6, e), chapter3_limits(m, e)
            for p in lo:
                self.assertAlmostEqual(lo[p], hi[p], delta=0.02,
                                       msg=f"разрыв предела {p} при {m} т, {e} дв.")

    def test_foreign_file_refused_with_headers(self):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["foo", "bar", "baz"])
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(ValueError) as cm:
            E.parse(buf.getvalue(), {"source_id": "easa_noise"})
        self.assertIn("foo", str(cm.exception), "отказ обязан показать, что увидел")


if __name__ == "__main__":
    unittest.main(verbosity=2)
