"""Таблица дозаполнения каталога: шаблон → заполненные строки → JSON, который принимает парсер."""

from __future__ import annotations

import importlib.util
import json
import pathlib
import shutil
import sys
import tempfile
import unittest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
spec = importlib.util.spec_from_file_location("catalog_fill", ROOT / "tools" / "catalog_fill.py")
CF = importlib.util.module_from_spec(spec)
spec.loader.exec_module(CF)

from flightcostapp.parsers.aircraft_user import _one   # noqa: E402
from flightcostapp.seo import DEV_PREFIX               # noqa: E402


class CatalogFill(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        d = pathlib.Path(self.tmp.name)
        self.cat = d / "catalog"
        shutil.copytree(ROOT / "data" / "aircraft" / "catalog", self.cat)
        self._old, CF.CATALOG = CF.CATALOG, self.cat
        self.xlsx = d / "t.xlsx"
        CF.template(self.xlsx)

    def tearDown(self):
        CF.CATALOG = self._old
        self.tmp.cleanup()

    def fill(self, sheet, icao, values, nth=0):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        ws = wb[sheet]
        head = [c.value for c in ws[1]]
        rows = [r for r in ws.iter_rows(min_row=2) if r[0].value == icao]
        for k, v in values.items():
            rows[nth][head.index(k)].value = v
        wb.save(self.xlsx)

    def doc(self, icao):
        return json.loads((self.cat / f"{icao}.json").read_text(encoding="utf-8"))

    def test_template_lists_only_gaps(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        burn = {r[0] for r in wb["Расход"].iter_rows(min_row=2, values_only=True)}
        self.assertIn("AT76", burn)
        self.assertIn("MS21", burn)
        self.assertNotIn("BCS3", burn, "у A220-300 якорь есть")
        at = [r for r in wb["Расход"].iter_rows(min_row=2, values_only=True) if r[0] == "AT76"]
        self.assertEqual([r[3] for r in at], [200, 400, 600], "турбовинтовой — три дистанции")

    def test_apply_anchor_and_passport(self):
        self.fill("Расход", "MS21", {"Дистанция, nm": 700, "Топливо на рейс, кг": 4300, "Пассажиров": 160,
                                     "Источник (издатель, документ, год)": "сводка по трём рейсам, 2026",
                                     "Вид данных": "оценка разработчика"})
        self.fill("Паспорт", "A338", {"Значение": 132.0, "Источник (издатель, документ, год)": "Airbus ACAP 2024",
                                      "Адрес": "https://aircraft.airbus.com/x.pdf", "Вид данных": "документ"})
        changed = CF.apply(self.xlsx)
        self.assertEqual(set(changed), {"MS21", "A338"})
        b = self.doc("MS21")["burn"]
        self.assertEqual((b["anchor_trip_nm"], b["anchor_trip_kg"], b["anchor_trip_pax"]), (700, 4300, 160))
        self.assertTrue(b["anchor_source"].startswith(DEV_PREFIX))
        a = self.doc("A338")
        self.assertEqual(a["fields"]["oew_t"], 132.0)
        self.assertNotIn("oew_t", (a["meta"].get("assumed") or {}))
        facts = _one(a, {"source_id": "aircraft_catalog"})
        self.assertFalse([f for f in facts if f.key.endswith("/calc_gaps")], "OEW закрыл пробел")

    def test_turboprop_table_without_analog(self):
        for i, (nm, kg) in enumerate(((200, 420), (400, 760), (600, 1100))):
            self.fill("Расход", "AT76", {"Топливо на рейс, кг": kg, "Пассажиров": 60,
                                         "Источник (издатель, документ, год)": "ATR, брошюра 2020",
                                         "Адрес": "https://www.atr-aircraft.com/x.pdf",
                                         "Вид данных": "документ"}, nth=i)
        CF.apply(self.xlsx)
        d = self.doc("AT76")
        self.assertEqual([x["nm"] for x in d["burn_table"]], [200, 400, 600])
        self.assertNotIn("burn", d, "аналог турбовинтовому не назначается сам")

    def test_document_needs_address_and_nothing_written(self):
        self.fill("Паспорт", "A338", {"Значение": 132.0, "Источник (издатель, документ, год)": "Airbus",
                                      "Вид данных": "документ"})
        before = (self.cat / "A338.json").read_text(encoding="utf-8")
        with self.assertRaises(SystemExit):
            CF.apply(self.xlsx)
        self.assertEqual((self.cat / "A338.json").read_text(encoding="utf-8"), before)


if __name__ == "__main__":
    unittest.main()
